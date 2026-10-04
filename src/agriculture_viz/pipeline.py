"""Provider adapters, immutable vintages, and reproducible frontend exports."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import tempfile
import urllib.request
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, cast

import polars as pl
import pycountry
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
PUBLIC = ROOT / "web" / "public" / "data"


def registry() -> dict[str, Any]:
    return cast(dict[str, Any], yaml.safe_load((ROOT / "sources.yaml").read_text()))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False))
    temp.replace(path)


def fetch(url: str, max_bytes: int = 100_000_000) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": "AgricultureRadar/0.1"})
    with urllib.request.urlopen(request, timeout=90) as response:
        body = cast(bytes, response.read(max_bytes + 1))
    if len(body) > max_bytes:
        raise ValueError(f"Provider response exceeds {max_bytes} bytes")
    return body


def download_inputs(config: dict[str, Any], cached: bool) -> dict[str, Any]:
    raw = DATA / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    providers = config["providers"]

    def faostat() -> tuple[str, dict[str, Any]]:
        cat_path = raw / "faostat-catalogue.json"
        archive = raw / "faostat.zip"
        if not cached:
            catalogue = json.loads(fetch(providers["faostat"]["catalogue_url"], 3_000_000))
            dataset = next(
                d
                for d in catalogue["Datasets"]["Dataset"]
                if d["DatasetCode"] == providers["faostat"]["dataset_code"]
            )
            archive.write_bytes(fetch(dataset["FileLocation"]))
            write_json(cat_path, dataset)
        dataset = json.loads(cat_path.read_text())
        return "faostat", record_release(
            "faostat",
            archive,
            dataset["FileLocation"],
            dataset.get("DateUpdate", "").split("T")[0] or None,
            providers["faostat"],
        )

    def worldbank() -> tuple[str, dict[str, Any]]:
        path = raw / "worldbank.xlsx"
        meta_path = raw / "worldbank-download.json"
        if not cached:
            html = fetch(providers["worldbank"]["source_url"], 3_000_000).decode()
            links = re.findall(r"href=[\"']([^\"']+)[\"']", html)
            url = next(u for u in links if providers["worldbank"]["download_match"] in u)
            path.write_bytes(fetch(url, 5_000_000))
            write_json(meta_path, {"download_url": url})
        url = json.loads(meta_path.read_text())["download_url"]
        workbook = load_workbook(path, read_only=True, data_only=True)
        published = None
        for row in list(workbook["Monthly Prices"].iter_rows(values_only=True))[:5]:
            if row[0] and str(row[0]).startswith("Updated on "):
                try:
                    published = (
                        datetime.strptime(str(row[0])[11:], "%B %d, %Y")
                        .replace(tzinfo=UTC)
                        .date()
                        .isoformat()
                    )
                except ValueError:
                    pass
        workbook.close()
        return "worldbank", record_release(
            "worldbank", path, url, published, providers["worldbank"]
        )

    def geography() -> tuple[str, dict[str, Any]]:
        path = raw / "geography.geojson"
        provider = providers["geography"]
        if not cached:
            path.write_bytes(fetch(provider["download_url"], 5_000_000))
        return "geography", record_release(
            "geography", path, provider["download_url"], None, provider
        )

    with ThreadPoolExecutor(max_workers=3) as pool:
        return dict(pool.map(lambda function: function(), (faostat, worldbank, geography)))


def record_release(
    provider: str, path: Path, url: str, published: str | None, config: dict[str, Any]
) -> dict[str, Any]:
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    release_id = f"{provider}-{checksum[:16]}"
    meta_path = DATA / "raw" / f"{release_id}.json"
    if meta_path.exists():
        result = cast(dict[str, Any], json.loads(meta_path.read_text()))
        if result["raw_file_hash"] != checksum:
            raise ValueError("Release hash collision")
        return result
    retained = DATA / "raw" / f"{release_id}{path.suffix}"
    shutil.copyfile(path, retained)
    first_seen = datetime.fromtimestamp(path.stat().st_mtime, UTC).isoformat()
    result = {
        "provider_release_id": release_id,
        "provider": provider,
        "provider_name": config["name"],
        "published_at": published,
        "publication_precision": "day" if published else None,
        "first_seen_at": first_seen,
        "retrieved_at": first_seen,
        "source_url": config["source_url"],
        "download_url": url,
        "raw_file_hash": checksum,
        "raw_path": str(retained.relative_to(ROOT)),
        "license": config["license"],
        "license_url": config.get("license_url", config["source_url"]),
    }
    write_json(meta_path, result)
    return result


def normalize_value(metric: str, unit: str, value: float) -> float:
    if not math.isfinite(value) or value < 0:
        raise ValueError("Nonfinite or negative crop observation")
    if metric == "yield":
        factors = {"kg/ha": 0.001, "hg/ha": 0.0001, "t/ha": 1.0}
        if unit not in factors:
            raise ValueError(f"Unsupported yield unit: {unit}")
        return value * factors[unit]
    expected = {"production": "t", "area": "ha"}
    if unit != expected[metric]:
        raise ValueError(f"Unexpected {metric} unit: {unit}")
    return value


def country_for_m49(value: str) -> Any | None:
    return pycountry.countries.get(numeric=value.lstrip("'").zfill(3))


def parse_production(
    path: Path, config: dict[str, Any], release_id: str
) -> tuple[pl.DataFrame, dict[str, Any]]:
    crops = {str(c["faostat_item_code"]): key for key, c in config["commodities"].items()}
    metrics = config["metrics"]
    with zipfile.ZipFile(path) as archive:
        members = [n for n in archive.namelist() if "Normalized" in n and n.endswith(".csv")]
        if len(members) != 1:
            raise ValueError("Expected one normalized FAOSTAT CSV")
        with archive.open(members[0]) as source, tempfile.NamedTemporaryFile(suffix=".csv") as tmp:
            shutil.copyfileobj(source, tmp)
            tmp.flush()
            frame = (
                pl.scan_csv(tmp.name, infer_schema=False)
                .filter(
                    pl.col("Item Code").is_in(list(crops)),
                    pl.col("Element Code").is_in(list(metrics)),
                )
                .collect()
            )
    countries: dict[str, Any] = {}
    rows = []
    for row in frame.iter_rows(named=True):
        country = country_for_m49(row["Area Code (M49)"])
        if country is None:
            continue  # Provider regions, China aggregate, and retired geographies are excluded.
        iso = country.alpha_3
        if iso in countries and countries[iso]["provider_area_code"] != row["Area Code"]:
            raise ValueError(f"Ambiguous geography crosswalk: {iso}")
        countries[iso] = {
            "id": iso,
            "name": country.name,
            "provider_name": row["Area"],
            "m49": country.numeric,
            "provider_area_code": row["Area Code"],
        }
        if not row["Value"] or row["Flag"] == "M":
            continue
        crop = crops[row["Item Code"]]
        metric = metrics[row["Element Code"]]["key"]
        unit = row["Unit"]
        if unit not in metrics[row["Element Code"]]["provider_units"]:
            raise ValueError(f"Unexpected provider unit {unit}")
        value = normalize_value(metric, unit, float(row["Value"]))
        year = int(row["Year"])
        rows.append(
            {
                "unique_id": f"{iso}.{crop}.{metric}",
                "geography_id": iso,
                "commodity_id": crop,
                "metric": metric,
                "unit": metrics[row["Element Code"]]["unit"],
                "year": year,
                "ds": date(year, 1, 1),
                "period_start": f"{year}-01-01",
                "period_end": f"{year}-12-31",
                "crop_year": None,
                "y": value,
                "status_flag": row["Flag"],
                "source_value": float(row["Value"]),
                "source_unit": unit,
                "provider_release_id": release_id,
            }
        )
    result = pl.DataFrame(rows)
    if result.select(["unique_id", "ds"]).is_duplicated().any():
        raise ValueError("Duplicate normalized production observations")
    return result.sort(["unique_id", "ds"]), countries


def parse_prices(
    path: Path, config: dict[str, Any], release_id: str
) -> tuple[pl.DataFrame, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["Monthly Prices"].iter_rows(values_only=True))
    header = next(row for row in rows if any(str(x).strip() == "Maize" for x in row))
    unit_row = rows[rows.index(header) + 1]
    columns: dict[str, int] = {}
    for crop, c in config["commodities"].items():
        columns[crop] = next(
            i for i, value in enumerate(header) if str(value).strip() == c["benchmark_column"]
        )
        if str(unit_row[columns[crop]]).strip() != "($/mt)":
            raise ValueError("Expected USD per metric tonne benchmark prices")
    output = []
    for row in rows:
        if not isinstance(row[0], str) or not re.fullmatch(r"\d{4}M\d{2}", row[0]):
            continue
        year, month = int(row[0][:4]), int(row[0][5:])
        for crop, column in columns.items():
            value = row[column]
            if not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
                continue
            output.append(
                {
                    "unique_id": f"GLOBAL.{crop}.price",
                    "commodity_id": crop,
                    "ds": date(year, month, 1),
                    "y": float(value),
                    "unit": "USD / tonne",
                    "status_flag": "published",
                    "provider_release_id": release_id,
                }
            )
    descriptions: dict[str, Any] = {}
    for row in workbook["Description"].iter_rows(values_only=True):
        text = next((str(v) for v in row[1:4] if v and len(str(v)) > 40), None)
        if text:
            for crop, c in config["commodities"].items():
                if (
                    crop == "wheat"
                    and "hard red winter" in text.lower()
                    or crop == "maize"
                    and text.lower().startswith(("maize", "corn"))
                    or crop == "rice"
                    and text.lower().startswith("rice (thailand), 5%")
                ):
                    descriptions[crop] = text
    workbook.close()
    result = pl.DataFrame(output).sort(["unique_id", "ds"])
    if result.select(["unique_id", "ds"]).is_duplicated().any():
        raise ValueError("Duplicate benchmark observations")
    return result, descriptions


def simplify_geography(path: Path, countries: dict[str, Any]) -> dict[str, Any]:
    data = json.loads(path.read_text())
    features = []
    for feature in data["features"]:
        p = feature["properties"]
        if p.get("ADMIN") == "Antarctica":
            continue
        country = country_for_m49(str(p.get("UN_A3", "")))
        iso = country.alpha_3 if country else p.get("ISO_A3_EH", p.get("ADM0_A3"))
        name = countries.get(iso, {}).get("name", p.get("ADMIN", p.get("NAME")))
        features.append(
            {
                "type": "Feature",
                "id": iso,
                "properties": {"id": iso, "name": name},
                "geometry": feature["geometry"],
            }
        )
    return {"type": "FeatureCollection", "features": features}


def export(vintage: Path) -> dict[str, Any]:
    """Export only from a selected immutable vintage; no provider access."""
    meta = json.loads((vintage / "metadata.json").read_text())
    production = pl.read_parquet(vintage / "production.parquet")
    prices = pl.read_parquet(vintage / "prices.parquet")
    PUBLIC.mkdir(parents=True, exist_ok=True)
    crop_metadata = {}
    for crop, config in meta["commodities"].items():
        frame = production.filter(pl.col("commodity_id") == crop)
        countries: dict[str, Any] = {}
        for row in frame.iter_rows(named=True):
            iso = row["geography_id"]
            countries.setdefault(iso, {})
            year = str(row["year"])
            countries[iso].setdefault(year, {})
            countries[iso][year][row["metric"]] = row["y"]
            countries[iso][year].setdefault("flags", {})[row["metric"]] = row["status_flag"]
        observations = prices.filter(pl.col("commodity_id") == crop).select("ds", "y").to_dicts()
        points = [{"date": row["ds"].isoformat(), "value": row["y"]} for row in observations]
        years = sorted(frame["year"].unique().to_list())
        output = {"commodity": crop, "years": years, "countries": countries, "prices": points}
        write_json(PUBLIC / f"{crop}.json", output)
        crop_metadata[crop] = {
            **config,
            "years": years,
            "latest_year": max(years),
            "country_count": len(countries),
            "latest_price_period": points[-1]["date"] if points else None,
            "benchmark_description": meta["benchmark_descriptions"].get(crop),
        }
        download_dir = PUBLIC / "downloads"
        download_dir.mkdir(exist_ok=True)
        export_columns = [
            "geography_id",
            "commodity_id",
            "metric",
            "unit",
            "year",
            "y",
            "status_flag",
            "source_value",
            "source_unit",
            "provider_release_id",
        ]
        frame.select(export_columns).write_csv(download_dir / f"{crop}-production.csv")
        prices.filter(pl.col("commodity_id") == crop).write_csv(download_dir / f"{crop}-prices.csv")
    manifest = {
        "schema_version": 1,
        "vintage": vintage.name,
        "generated_at": meta["generated_at"],
        "countries": meta["countries"],
        "commodities": crop_metadata,
        "releases": meta["releases"],
        "flags": meta["flags"],
        "coverage_definition": "Current ISO country/territory crosswalk; provider regional aggregates and retired geographies excluded.",
        "world_share_definition": "Country production / sum of covered country and territory production in the same crop and year. Missing observations are not zero.",
        "change_definition": "Percentage change from mean production of five calendar years immediately preceding the selected year; all five observations required.",
    }
    write_json(PUBLIC / "manifest.json", manifest)
    shutil.copyfile(vintage / "geography.geojson", PUBLIC / "geography.geojson")
    forecast_path = DATA / "forecast" / "latest.json"
    if forecast_path.exists():
        shutil.copyfile(forecast_path, PUBLIC / "forecast.json")
    return manifest


def refresh(*, cached: bool = False) -> Path:
    config = registry()
    releases = download_inputs(config, cached)
    raw_paths = {key: ROOT / value["raw_path"] for key, value in releases.items()}
    production, countries = parse_production(
        raw_paths["faostat"], config, releases["faostat"]["provider_release_id"]
    )
    prices, descriptions = parse_prices(
        raw_paths["worldbank"], config, releases["worldbank"]["provider_release_id"]
    )
    now = datetime.now(UTC)
    vintage = DATA / "vintages" / now.strftime("%Y-%m-%dT%H%M%S.%fZ")
    vintage.mkdir(parents=True, exist_ok=False)
    production.write_parquet(vintage / "production.parquet")
    prices.write_parquet(vintage / "prices.parquet")
    write_json(vintage / "geography.geojson", simplify_geography(raw_paths["geography"], countries))
    write_json(
        vintage / "metadata.json",
        {
            "generated_at": now.isoformat(),
            "countries": countries,
            "commodities": config["commodities"],
            "releases": releases,
            "benchmark_descriptions": descriptions,
            "flags": config["flags"],
        },
    )
    export(vintage)
    write_json(DATA / "latest.json", {"vintage": vintage.name})
    return vintage


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cached", action="store_true", help="Use retained raw inputs")
    parser.add_argument(
        "--export-vintage", help="Re-export one retained vintage, without network access"
    )
    args = parser.parse_args()
    if args.export_vintage:
        vintage = DATA / "vintages" / args.export_vintage
        if vintage.resolve().parent != (DATA / "vintages").resolve():
            raise ValueError("Invalid vintage path")
        export(vintage)
    else:
        vintage = refresh(cached=args.cached)
    print(json.dumps({"vintage": vintage.name, "export": str(PUBLIC)}))


if __name__ == "__main__":
    main()
